"""كشفُ الحصص في Excel — بأدوات المنصّة المشتركة لا بنمطٍ خاصّ.

ترويسةُ المنصّة الواحدة (`add_excel_title_rows`)، وجدولُها (`excel_table_styles`)،
وحمايةُ الورقة للقراءة وإعدادُ الطباعة من `ExcelService`: A4 أفقيٌّ بعرض صفحةٍ
واحدة وترويسةٍ تتكرّر. وألوانُ الحال من `core.brand` — لا لونَ يُكتب هنا.
"""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Side
from openpyxl.utils import get_column_letter

from core import brand
from core.export_utils import add_excel_title_rows, excel_table_styles, xl_fill, xl_font
from reports.services import ExcelService

from .register import LETTER, UNRECORDED, WHERE_CODE

HEADER_ROW = 4
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
START = Alignment(horizontal="right", vertical="center", wrap_text=True)

#: لونُ الخانة بحالها: خلفيّةٌ فاتحةٌ ونصٌّ داكنٌ يُقرأ مطبوعاً بالأبيض والأسود.
STATUS_STYLE = {
    "present": (brand.STATUS_SUCCESS_BG, brand.STATUS_SUCCESS_DARK),
    "absent": (brand.STATUS_DANGER_BG, brand.STATUS_DANGER_DARK),
    "late": (brand.STATUS_WARNING_BG, brand.STATUS_WARNING_DARK),
}
LEGEND = (
    " · ".join(
        f"{letter} {label}"
        for letter, label in zip(
            LETTER.values(), ("حاضر", "غائب", "متأخّر (والرقمُ دقائقُه)"), strict=True
        )
    )
    + f" · {UNRECORDED} لم تُرصد  —  أين الطالب: "
    + " · ".join(
        f"{code} {label}"
        for code, label in zip(
            WHERE_CODE.values(),
            ("عيادة", "نشاط", "خرج بإذن", "خرج دون إذن", "استئذان مبكّر", "البوّابة"),
            strict=True,
        )
    )
)


def _merge_line(ws, row: int, num_cols: int, value: str, *, font=None, fill=None, align=START):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=num_cols)
    cell = ws.cell(row=row, column=1, value=value)
    cell.font = font or xl_font(brand.TEXT_SECONDARY, size=9)
    cell.alignment = align
    if fill is not None:
        cell.fill = fill
    return cell


def _stamp(register) -> str:
    when = register.as_of.strftime("%H:%M")
    return f"كشفٌ جزئيٌّ حتى {when}" if register.is_partial else f"كشفٌ كامل — صدر {when}"


def _signatures(ws, row: int, num_cols: int, signatories) -> int:
    """ثلاثُ خاناتِ توقيعٍ متجاورة: المسمّى، ثمّ الاسم، ثمّ سطرُ التوقيع."""
    row += 1
    third = max(1, num_cols // 3)
    line = Border(bottom=Side(style="thin", color=brand.excel(brand.TEXT_PRIMARY)))
    for index, person in enumerate(signatories):
        first = 1 + index * third
        last = num_cols if index == len(signatories) - 1 else first + third - 1
        for offset, (value, font, height) in enumerate(
            (
                (person.title, xl_font(brand.MAROON, size=10, bold=True), 20),
                (person.name or " ", xl_font(brand.TEXT_PRIMARY, size=10), 18),
                ("", xl_font(), 30),
            )
        ):
            r = row + offset
            ws.merge_cells(start_row=r, start_column=first, end_row=r, end_column=last)
            cell = ws.cell(row=r, column=first, value=value)
            cell.font = font
            cell.alignment = CENTER
            ws.row_dimensions[r].height = height
            if offset == 2:
                for col in range(first, last + 1):
                    ws.cell(row=r, column=col).border = line
    return row + 3


def write_section_sheet(
    ws, register, school_name: str, exported_by: str, orientation: str = "landscape"
) -> None:
    """ورقةُ شعبةٍ واحدة: الطلابُ صفوفاً، والحصصُ أعمدة، وذيلٌ لكلّ حصّة، ثمّ التواقيع."""
    ws.sheet_view.rightToLeft = True
    columns = register.columns
    num_cols = 2 + len(columns) + 3
    table = excel_table_styles()

    wing = register.wing.name if register.wing else "خارج الأجنحة"
    add_excel_title_rows(
        ws,
        num_cols,
        f"{school_name}  —  وزارة التربية والتعليم والتعليم العالي",
        f"{register.title}  —  {register.day:%Y/%m/%d}",
        f"{wing}  |  {_stamp(register)}  |  صدر بواسطة: {exported_by}",
    )

    headers = [("#", 5), ("الطالب", 34)]
    headers += [(f"ح{c.number}\n{c.start:%H:%M}–{c.end:%H:%M}", 11) for c in columns]
    headers += [("حصص الغياب", 10), ("مرّات التأخّر", 10), ("دقائق التأخّر", 10)]
    ExcelService._add_header_row(
        ws,
        {
            "header_font": table.header_font,
            "header_fill": table.header_fill,
            "header_align": table.header_align,
            "thin_border": table.border,
        },
        HEADER_ROW,
        headers,
    )
    ws.row_dimensions[HEADER_ROW].height = 34

    row = HEADER_ROW
    for student in register.rows:
        row += 1
        values = [student.number, student.name]
        values += [mark.text for mark in student.marks]
        values += [student.absent_periods, student.late_count, student.late_minutes]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = table.border
            cell.alignment = START if col == 2 else CENTER
            cell.font = table.cell_font
        for offset, mark in enumerate(student.marks):
            if mark.status in STATUS_STYLE:
                bg, fg = STATUS_STYLE[mark.status]
                cell = ws.cell(row=row, column=3 + offset)
                cell.fill = xl_fill(bg)
                cell.font = xl_font(fg, bold=True)
        ws.row_dimensions[row].height = 20

    # ذيلُ الحصص: الحالُ، ووقتُ أوّل تثبيت، ومن ثبّت، والعدد.
    footer = (
        ("الحال", lambda c: c.label),
        (
            "أوّل تثبيت",
            lambda c: c.first_confirmed_at.strftime("%H:%M:%S") if c.is_confirmed else "—",
        ),
        ("ثبّتها", lambda c: c.confirmed_by or "—"),
        (
            "حاضر / غائب / متأخّر",
            lambda c: f"{c.present} / {c.absent} / {c.late}" if c.is_confirmed else "—",
        ),
    )
    for label, value_of in footer:
        row += 1
        head = ws.cell(row=row, column=2, value=label)
        head.font = xl_font(brand.MAROON, size=9, bold=True)
        head.alignment = START
        head.fill = xl_fill(brand.MAROON_BG)
        for offset, column in enumerate(columns):
            cell = ws.cell(row=row, column=3 + offset, value=value_of(column))
            cell.font = xl_font(
                brand.STATUS_WARNING_DARK
                if column.status == "confirmed_late"
                else brand.TEXT_SECONDARY,
                size=8,
                bold=column.status == "confirmed_late",
            )
            cell.alignment = CENTER
            cell.fill = xl_fill(brand.MAROON_BG)
            cell.border = table.border
        ws.row_dimensions[row].height = 26 if label == "ثبّتها" else 18

    row += 1
    _merge_line(ws, row, num_cols, LEGEND, font=xl_font(brand.TEXT_MUTED, size=8))
    ws.row_dimensions[row].height = 28
    row = _signatures(ws, row, num_cols, register.signatories)

    ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=3)
    _setup_page(ws, num_cols, row, school_name, orientation)
    ExcelService._apply_protection(ws, num_cols)


def write_wing_summary(
    ws, register, school_name: str, exported_by: str, orientation: str = "landscape"
) -> None:
    """مصفوفةُ الجناح: الشُّعبُ صفوفاً وحصصُ اليوم أعمدة — وقتُ التثبيت وعددُ الغائبين."""
    ws.sheet_view.rightToLeft = True
    table = excel_table_styles()
    width = register.width
    num_cols = 1 + width + 3
    add_excel_title_rows(
        ws,
        num_cols,
        f"{school_name}  —  وزارة التربية والتعليم والتعليم العالي",
        f"{register.title}  —  {register.day:%Y/%m/%d}",
        f"ملخّصُ رصد الحصص  |  {_stamp(register)}  |  صدر بواسطة: {exported_by}",
    )
    headers = [("الشعبة", 14)] + [(f"ح{n}", 16) for n in range(1, width + 1)]
    headers += [("حصصٌ مثبّتة", 12), ("ثُبّتت متأخّرة", 12), ("لم تُرصد", 12)]
    ExcelService._add_header_row(
        ws,
        {
            "header_font": table.header_font,
            "header_fill": table.header_fill,
            "header_align": table.header_align,
            "thin_border": table.border,
        },
        HEADER_ROW,
        headers,
    )

    row = HEADER_ROW
    for section, cells in register.matrix:
        row += 1
        head = ws.cell(row=row, column=1, value=section.class_group.short_code)
        head.font = xl_font(brand.TEXT_PRIMARY, bold=True)
        head.alignment = CENTER
        head.border = table.border
        for offset, column in enumerate(cells):
            cell = ws.cell(row=row, column=2 + offset)
            cell.border = table.border
            cell.alignment = CENTER
            if column is None:
                cell.value = ""
                continue
            if column.is_confirmed:
                late = column.status == "confirmed_late"
                cell.value = f"{'متأخّرة ' if late else ''}{column.first_confirmed_at:%H:%M}\nغ {column.absent} · م {column.late}"
                bg, fg = (
                    (brand.STATUS_WARNING_BG, brand.STATUS_WARNING_DARK)
                    if late
                    else (brand.STATUS_SUCCESS_BG, brand.STATUS_SUCCESS_DARK)
                )
            else:
                cell.value = column.label
                bg, fg = (
                    (brand.STATUS_DANGER_BG, brand.STATUS_DANGER_DARK)
                    if column.status == "missed"
                    else (brand.SURFACE_ALT, brand.TEXT_MUTED)
                )
            cell.fill = xl_fill(bg)
            cell.font = xl_font(fg, size=9, bold=True)
        confirmed = sum(c.is_confirmed for c in section.columns)
        late = section.late_confirmations
        missed = sum(c.status == "missed" for c in section.columns)
        for offset, value in enumerate((confirmed, late, missed)):
            cell = ws.cell(row=row, column=2 + width + offset, value=value)
            cell.border = table.border
            cell.alignment = CENTER
            cell.font = xl_font(bold=True)
        ws.row_dimensions[row].height = 34

    row += 1
    _merge_line(
        ws,
        row,
        num_cols,
        "في الخانة وقتُ أوّل تثبيتٍ وعددُ الغائبين (غ) والمتأخّرين (م). «متأخّرة»: ثُبّتت بعد "
        "خمس دقائق من نهاية الحصّة. «لم تُرصد»: انقضت مهلتُها ولم تُثبَّت.",
        font=xl_font(brand.TEXT_MUTED, size=8),
    )
    row = _signatures(ws, row, num_cols, register.signatories)
    _setup_page(ws, num_cols, row, school_name, orientation)
    ExcelService._apply_protection(ws, num_cols)


ORIENTATIONS = ("landscape", "portrait")


def _setup_page(ws, num_cols: int, last_row: int, school_name: str, orientation: str) -> None:
    """A4 بالاتّجاه المختار. والعموديُّ **صفحةٌ واحدةٌ لكلّ ورقة** (قرارُ 2026-09-13):
    يُصغَّر عرضاً وطولاً معاً، فلا تنقلب التواقيعُ إلى صفحةٍ ثانية."""
    orientation = orientation if orientation in ORIENTATIONS else "landscape"
    ExcelService._setup_print(
        ws,
        num_cols,
        last_row - HEADER_ROW,
        paper="a4",
        orientation=orientation,
        header_text=school_name,
    )
    ws.print_area = f"A1:{get_column_letter(num_cols)}{last_row}"
    if orientation == "portrait":
        ws.page_setup.fitToHeight = 1


def _sheet_title(text: str) -> str:
    """اسمُ الورقة: 31 حرفاً بلا المحارف التي يرفضها Excel."""
    for bad in "[]:*?/\\":
        text = text.replace(bad, "-")
    return text[:31]


def section_workbook(register, school_name: str, exported_by: str, orientation="landscape"):
    wb, ws, _ = ExcelService._make_workbook(_sheet_title(register.class_group.short_code))
    write_section_sheet(ws, register, school_name, exported_by, orientation)
    return wb


def wing_workbook(register, school_name: str, exported_by: str, orientation="landscape"):
    wb, ws, _ = ExcelService._make_workbook("ملخّص الجناح")
    write_wing_summary(ws, register, school_name, exported_by, orientation)
    for section in register.sections:
        sheet = wb.create_sheet(_sheet_title(section.class_group.short_code))
        write_section_sheet(sheet, section, school_name, exported_by, orientation)
    return wb
