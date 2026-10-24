"""التقريرُ الشهريّ للعيادة Excel (W-20261020-032): ترتيبُ الخدمات، المجاميع، الإخفاءُ، النطاقُ، القدرةُ المستقلّة، وخلوُّه من الهويّات."""

import datetime as dt
from io import BytesIO

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from clinic import choices
from clinic import export as ce
from clinic.models import ClinicVisit
from core.capabilities.access import has_capability
from core.exports import registry

pytestmark = pytest.mark.django_db


def _visit(school, student, nurse, when, **kw):
    v = ClinicVisit.objects.create(
        school=school, student=student, nurse=nurse, reason="سرٌّ سريريّ", **kw
    )
    ClinicVisit.objects.filter(pk=v.pk).update(visit_date=when)
    return v


def _in_month(day=10):
    return timezone.make_aware(dt.datetime(2026, 9, day, 9, 0))


def test_service_rows_follow_the_owner_order_and_totals_are_exact(school, student_user, nurse_user):
    student_user.nationality = "قطري"
    student_user.save()
    for _ in range(7):
        _visit(school, student_user, nurse_user, _in_month(), service_code="bandage")
    counts = ce.monthly_counts(school, 2026, 9)
    rows = ce.service_rows(counts)
    assert [r[1] for r in rows] == [label for _c, label, _g in choices.SERVICES]
    first = rows[0]
    assert first[2:] == (7, 0, 7)
    assert counts["total"] == 7


def test_small_counts_are_suppressed_but_service_total_stays_exact(
    school, student_user, nurse_user
):
    for _ in range(3):
        _visit(school, student_user, nurse_user, _in_month(), service_code="insulin")
    counts = ce.monthly_counts(school, 2026, 9)
    row = next(r for r in ce.service_rows(counts) if r[1] == "إنسولين")
    assert row[3] == ce.SUPPRESSED and row[4] == 3


def test_other_months_and_other_schools_are_excluded(school, student_user, nurse_user):
    _visit(
        school,
        student_user,
        nurse_user,
        timezone.make_aware(dt.datetime(2026, 8, 31, 23)),
        service_code="bandage",
    )
    _visit(
        school,
        student_user,
        nurse_user,
        timezone.make_aware(dt.datetime(2026, 10, 1, 1)),
        service_code="bandage",
    )
    assert ce.monthly_counts(school, 2026, 9)["total"] == 0


def test_workbook_has_two_sheets_and_no_identity_or_free_text(school, student_user, nurse_user):
    _visit(
        school,
        student_user,
        nurse_user,
        _in_month(),
        service_code="bandage",
        complaint_code="headache",
    )
    wb = load_workbook(BytesIO(ce.build_workbook(school, 2026, 9)))
    assert wb.sheetnames == ["السجل", "اللوحة"]
    text = " ".join(
        str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None
    )
    for forbidden in (student_user.full_name, nurse_user.full_name, "سرٌّ سريريّ"):
        assert forbidden not in text


def test_export_kind_is_registered_behind_its_own_capability():
    assert registry.get("clinic.monthly_xlsx").capability == "clinic.export"


def test_export_capability_is_independent_of_read_access(nurse_user, teacher_user):
    assert has_capability(nurse_user, "clinic.export")
    assert not has_capability(teacher_user, "clinic.export")


def test_export_capability_is_for_nurse_and_principal_only():
    """قرارُ المالك: تصديرُ العيادة للممرضة والمدير فقط — لا نائبَ إداريّ."""
    from core.capabilities import registry as caps

    assert caps()["clinic.export"].roles == {"nurse", "principal"}


def test_builder_returns_audit_row_count(school, student_user, nurse_user):
    _visit(school, student_user, nurse_user, _in_month(), service_code="bandage")
    res = ce.build_monthly_export(school, nurse_user, {"year": "2026", "month": "9"})
    assert res.rows == 1 and res.object_id == "2026-09" and not res.full_national_id


def test_pdf_kind_shares_the_export_capability():
    assert registry.get("clinic.monthly_pdf").capability == "clinic.export"


def test_monthly_pdf_is_landscape_with_confidentiality_and_manager_approval(
    school, student_user, nurse_user
):
    from django.template.loader import render_to_string
    from weasyprint import HTML

    _visit(school, student_user, nurse_user, _in_month(), service_code="bandage")
    res = ce.build_monthly_pdf(school, nurse_user, {"year": "2026", "month": "9"})
    assert res.content.startswith(b"%PDF") and res.content_type == "application/pdf"
    assert res.rows == 1 and res.object_id == "2026-09"

    html = render_to_string(
        "reports/clinic_monthly.html",
        {
            "school": school,
            "month_label": "سبتمبر 2026",
            "total": 1,
            "uncoded": 0,
            "service_rows": ce.service_rows(ce.monthly_counts(school, 2026, 9)),
            "dashboard_rows": [],
            "exported_by": "-",
            "exporter_role_ar": "-",
            "export_datetime": "-",
            "academic_year": "-",
        },
    )
    page = HTML(string=html).render().pages[0]
    assert page.width > page.height
    assert "سرّي" in html and "اعتماد مدير المدرسة" in html
    assert student_user.full_name not in html
