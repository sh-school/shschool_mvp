"""يومُ الكشف المفروض عند الاستيراد وكشفُ اليوم Excel (W-20261005-011، أكّده المالك 2026-10-05).

الأوّل اختياريٌّ يفرض يومَه على كلّ أسطر الملفّ، فلا يقع رصدُ الكشف في يومٍ غيرِ ما أراده السكرتير
ولا يفلت من نافذة الاستيراد؛ والثاني ملخّصٌ يوميٌّ بالحضور والانصراف لمن تراهم لوحةُ الرصد نفسُها.
"""

from __future__ import annotations

import io
from datetime import date

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from core.models import AuditLog
from staff_affairs.models import StaffAttendance
from tests.test_biometric_import import DAY, NOW, _csv, _line, _person, secretary  # noqa: F401

pytestmark = [pytest.mark.django_db, pytest.mark.pin_clock]

IMPORT = "staff_affairs:attendance_import"
COMMIT = "staff_affairs:attendance_import_commit"
DAILY = "staff_affairs:attendance_daily_xlsx"


@pytest.fixture(autouse=True)
def _clock(monkeypatch):
    import staff_affairs.attendance as attendance_module

    monkeypatch.setattr(
        attendance_module,
        "_now",
        lambda: timezone.localtime(timezone.make_aware(NOW)),
        raising=False,
    )
    # لوحةُ اليوم والنموذجُ يقرآن اليومَ من ساعة Django لا من ساعة الحضور.
    monkeypatch.setattr(timezone, "localdate", lambda *a, **k: NOW.date())


def _upload(client, *lines, day=None):
    data = {"file": SimpleUploadedFile("كشف.csv", _csv(*lines))}
    if day is not None:
        data["day"] = day
    return client.post(reverse(IMPORT), data)


class TestForcedDay:
    def test_without_a_day_each_line_keeps_its_own_date(self, client_as, school, secretary):  # noqa: F811
        _person(school, 4101)
        response = _upload(client_as(secretary), _line(4101, day="02/02/2026"))
        assert [i.row.day for i in response.context["plan"].items] == [DAY]

    def test_a_chosen_day_overrides_the_file_date(self, client_as, school, secretary):  # noqa: F811
        staff = _person(school, 4102)
        client = client_as(secretary)
        response = _upload(client, _line(4102, day="02/02/2026"), day="2026-02-05")
        assert [i.row.day for i in response.context["plan"].items] == [date(2026, 2, 5)]
        client.post(reverse(COMMIT), {"token": response.context["token"]})
        assert StaffAttendance.objects.filter(staff=staff, date=date(2026, 2, 5)).exists()
        assert not StaffAttendance.objects.filter(staff=staff, date=DAY).exists()

    def test_two_lines_of_one_person_collapse_to_a_duplicate_under_a_forced_day(
        self,
        client_as,
        school,
        secretary,  # noqa: F811
    ):
        _person(school, 4103)
        response = _upload(
            client_as(secretary),
            _line(4103, day="02/02/2026"),
            _line(4103, day="03/02/2026"),
            day="2026-02-05",
        )
        plan = response.context["plan"]
        assert plan.writable == 1
        assert [i.action for i in plan.items].count("invalid") == 1

    def test_a_future_day_is_refused_by_the_form(self, client_as, school, secretary):  # noqa: F811
        _person(school, 4104)
        response = _upload(client_as(secretary), _line(4104), day="2026-02-16")
        assert "plan" not in response.context
        assert "لم يأتِ" in response.content.decode()

    def test_a_day_older_than_the_window_is_refused_by_the_form(self, client_as, school, secretary):  # noqa: F811
        _person(school, 4105)
        response = _upload(client_as(secretary), _line(4105), day="2025-12-01")
        assert "plan" not in response.context
        assert "نافذة الاستيراد" in response.content.decode()

    def test_the_upload_page_offers_the_optional_day_field(self, client_as, secretary):  # noqa: F811
        html = client_as(secretary).get(reverse(IMPORT)).content.decode()
        assert 'name="day"' in html and "اختياريّ" in html


class TestDailySheet:
    def test_the_sheet_lists_everyone_with_status_and_times(self, client_as, school, secretary):  # noqa: F811
        staff = _person(school, 4201)
        _person(school, 4202)
        client = client_as(secretary)
        token = _upload(client, _line(4201, check_in="06:50AM", check_out="02:10PM")).context[
            "token"
        ]
        client.post(reverse(COMMIT), {"token": token})

        response = client.get(reverse(DAILY), {"date": "2026-02-02"})
        assert response.status_code == 200 and "spreadsheetml" in response["Content-Type"]
        sheet = load_workbook(io.BytesIO(response.content)).active
        rows = {r[0]: r for r in sheet.iter_rows(min_row=2, values_only=True)}
        assert rows["4201"][3] == "06:50" and rows["4201"][4] == "14:10"
        assert rows["4202"][2] == "لم يُرصد" and not rows["4202"][3]
        assert staff.full_name in rows["4201"][1]

    def test_the_sheet_never_carries_the_excuse_text(self, client_as, school, secretary):  # noqa: F811
        staff = _person(school, 4203)
        StaffAttendance.objects.create(
            school=school, staff=staff, date=DAY, status="absent",
            absence_type="sick", accepted_excuse="تقرير طبي سرّي",
        )  # fmt: skip
        response = client_as(secretary).get(reverse(DAILY), {"date": "2026-02-02"})
        assert "تقرير طبي".encode() not in response.content
        sheet = load_workbook(io.BytesIO(response.content)).active
        assert "تقرير طبي سرّي" not in {c for r in sheet.iter_rows(values_only=True) for c in r}

    def test_the_export_is_audited_without_names(self, client_as, school, secretary):  # noqa: F811
        _person(school, 4204)
        client_as(secretary).get(reverse(DAILY), {"date": "2026-02-02"})
        log = AuditLog.objects.filter(action="export").latest("timestamp")
        assert log.changes["kind"] == "staff_affairs.attendance_day_xlsx"
        assert log.changes["full_national_id"] is False

    def test_a_teacher_is_denied(self, client_as, school):
        assert client_as(_person(school, 4205)).get(reverse(DAILY)).status_code == 403

    def test_the_board_links_to_the_sheet_for_the_shown_day(self, client_as, secretary):  # noqa: F811
        html = (
            client_as(secretary)
            .get(reverse("staff_affairs:attendance_board"), {"date": "2026-02-02"})
            .content.decode()
        )
        assert f"{reverse(DAILY)}?date=2026-02-02" in html
