"""تصديرُ التقرير الشهري (سلوك وتعليم) بصفوفٍ فعليّة (W-20261008-001 · البند 4 — نصفُ التصدير).

كان يقرأ `r["behavior_points"]` والصفوفُ لا تحمله فيرمي KeyError عند أوّل صفّ؛ ونظامُ الخصم ملغى
فحُذف عمودُ «النقاط المخصومة» بدل اختلاق مفتاحٍ.
"""

import io

import pytest
from openpyxl import load_workbook

from reports.services import AcademicReportsExcel


@pytest.mark.django_db
def test_monthly_excel_with_rows_opens_and_has_no_points_column(school):
    data = {
        "period": "2026-04",
        "rows": [
            {
                "student_name": "طالب تجريبي",
                "student_id": "1",
                "quiz_count": 2,
                "quiz_avg": 80.0,
                "behavior_count": 1,
                "combined_score": 77.5,
            }
        ],
    }
    response = AcademicReportsExcel.monthly_behavior_academic_excel(data, school)
    ws = load_workbook(io.BytesIO(response.content)).active
    cells = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert "النقاط المخصومة" not in cells
    assert "عدد المخالفات" in cells and "التقييم المدمج" in cells
    assert "طالب تجريبي" in cells
