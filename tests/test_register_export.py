"""كشفُ الحصص مطبوعاً ومصدَّراً — كشفُ الشعبة وكشفُ الجناح.

قراراتُ 2026-09-13 التي تحرسها هذه الاختبارات:

- لا رقمَ شخصيّاً في الكشف، لا في الـPDF ولا في Excel.
- توقيعٌ واحد: مشرفُ الجناح وحدَه، لا النائبُ الإداريّ ولا المدير (D-385م، يُلغي قرارَ الثلاثة تواقيع).
- الكشفُ يقرأ رصدَ المعلّم من جدول الشعبة (`AttendanceEntry`) كما تقرؤه لوحةُ المشرف، بفرق 0 (W-20261011-011).
- الكشفُ الجزئيُّ يُطبع موسوماً «كشفٌ جزئيٌّ حتى HH:MM»، لا يُمنع.
- ملفُّ Excel محميٌّ للقراءة، ولكشف الجناح ورقةُ ملخّصٍ ثمّ ورقةٌ لكلّ شعبة.
- المشرفُ لجناحه وحدَه، كصلاحيّة الرصد.

ومعها ثلاثُ علل في أدوات التصدير المشتركة أُصلحت مع الكشف.
"""

import io

import openpyxl
import pytest
from django.urls import reverse

from core.export_utils import get_export_context
from core.models import Wing
from core.pdf_utils import _content_disposition
from tests.conftest import MembershipFactory, RoleFactory, UserFactory
from tests.test_period_register import (  # noqa: F401 — تجهيزاتُ الكشف نفسُها
    SUNDAY,
    _confirm,
    _periods,
    at,
    kids,
    klass,
    supervisor,
    teacher,
    year,
)
from wings.register import section_register, wing_register

pytestmark = pytest.mark.django_db


def _section_url(klass, fmt="html"):
    return (
        reverse("wings:section_register", args=[klass.id])
        + f"?date={SUNDAY.isoformat()}&format={fmt}"
    )


def _wing_url(code, fmt="html"):
    return reverse("wings:wing_register", args=[code]) + f"?date={SUNDAY.isoformat()}&format={fmt}"


# ══════════════════════════════════════════════════════════════════
# البيانات
# ══════════════════════════════════════════════════════════════════


class TestTheSectionRegister:
    def test_a_day_with_unconfirmed_periods_is_partial(
        self, school, klass, kids, teacher, supervisor
    ):
        periods = _periods(school, klass, teacher, 2)
        _confirm(klass, periods[0], {}, supervisor)

        assert section_register(klass, SUNDAY).is_partial
        _confirm(klass, periods[1], {}, supervisor)
        assert not section_register(klass, SUNDAY).is_partial

    def test_one_signature_the_wing_supervisor_only(self, school, klass, kids, teacher, supervisor):
        register = section_register(klass, SUNDAY)

        assert [(s.title, s.name) for s in register.signatories] == [("مشرف الجناح", "مشرف الجناح")]
        assert [s.title for s in wing_register(klass.wing, SUNDAY).signatories] == ["مشرف الجناح"]

    def test_no_vice_admin_or_principal_even_when_they_exist(
        self, school, klass, kids, teacher, supervisor
    ):
        for name, role, nid in (
            ("النائب الإداري", "vice_admin", "29300000094"),
            ("مدير المدرسة", "principal", "29300000095"),
        ):
            MembershipFactory(
                user=UserFactory(full_name=name, national_id=nid),
                school=school,
                role=RoleFactory(school=school, name=role),
            )

        titles = [s.title for s in section_register(klass, SUNDAY).signatories]

        assert titles == ["مشرف الجناح"]


def _enter(session, kid, by, status="absent", minutes=None, hour=8):
    """رصدُ المعلّم من جدول الشعبة — رأسٌ في دفتر الإدخال لا تثبيتُ مشرف."""
    from operations.models import AttendanceEntry

    return AttendanceEntry.objects.create(
        school=session.school,
        session=session,
        student=kid,
        status=status,
        tardiness_minutes=minutes,
        entered_by=by,
        entered_at=at(hour, 5),
    )


class TestTheRegisterReadsTheTeachersEntries:
    """الكشفُ كان يقرأ التثبيتَ القديم وحدَه فيعرض «لم تُرصد» لكلّ حصّةٍ رصدها المعلّم (W-20261011-011)."""

    def test_an_entered_period_is_not_shown_as_not_recorded(
        self, school, klass, kids, teacher, supervisor
    ):
        periods = _periods(school, klass, teacher, 2)
        for kid in kids:
            _enter(periods[0], kid, teacher, "absent" if kid == kids[0] else "present")

        register = section_register(klass, SUNDAY)

        first, second = register.columns
        assert first.is_confirmed and first.label == "رُصدت"
        assert first.status != "missed" and first.first_confirmed_at is not None
        assert not second.is_confirmed, "حصّةٌ بلا إدخالٍ تبقى بلا رصد"
        assert register.rows[0].marks[0].status == "absent"
        assert register.rows[1].marks[0].status == "present"
        assert register.rows[1].marks[1].status == "", "لا خانةَ مخترعة"

    def test_the_numbers_equal_the_supervisor_panels_with_zero_difference(
        self, school, klass, kids, teacher, supervisor
    ):
        from operations.attendance_selectors import entry_tallies_by_start

        periods = _periods(school, klass, teacher, 2)
        for index, period in enumerate(periods):
            _enter(period, kids[0], teacher, "absent")
            _enter(period, kids[1], teacher, "late", minutes=4)
            _enter(period, kids[2], teacher, "present")
            if index == 0:
                _enter(period, kids[3], teacher, "absent")

        panel = entry_tallies_by_start(klass, SUNDAY)
        register = section_register(klass, SUNDAY)

        for column in register.columns:
            tally = panel[column.start]
            assert (column.present, column.absent, column.late) == (
                tally["present"],
                tally["absent"],
                tally["late"],
            )
        assert (register.columns[0].absent, register.columns[1].absent) == (2, 1)

    def test_a_correction_counts_by_its_head_only(self, school, klass, kids, teacher, supervisor):
        from operations.models import AttendanceEntry

        period = _periods(school, klass, teacher, 1)[0]
        first = _enter(period, kids[0], teacher, "absent")
        AttendanceEntry.objects.create(
            school=school,
            session=period,
            student=kids[0],
            status="present",
            entered_by=teacher,
            entered_at=at(8, 20),
            supersedes=first,
            correction_reason="تصحيح",
        )

        register = section_register(klass, SUNDAY)

        assert register.rows[0].marks[0].status == "present"
        assert register.columns[0].absent == 0

    def test_a_day_with_no_entries_still_says_not_recorded(
        self, school, klass, kids, teacher, supervisor
    ):
        _periods(school, klass, teacher, 2)

        register = section_register(klass, SUNDAY)

        assert not any(c.is_confirmed for c in register.columns)
        assert all(m.status == "" for row in register.rows for m in row.marks)
        assert all((c.present, c.absent, c.late) == (0, 0, 0) for c in register.columns)

    def test_a_supervisor_confirmation_still_wins_over_entries(
        self, school, klass, kids, teacher, supervisor
    ):
        period = _periods(school, klass, teacher, 1)[0]
        _confirm(klass, period, {kids[0]: "absent"}, supervisor)

        column = section_register(klass, SUNDAY).columns[0]

        assert column.is_confirmed and column.label != "رُصدت"
        assert column.confirmed_by == "مشرف الجناح"


class TestTheWingRegister:
    def test_the_matrix_pads_a_shorter_day(self, school, year, klass, kids, teacher, supervisor):
        from tests.conftest import ClassGroupFactory

        short = ClassGroupFactory(
            school=school, grade="G7", section="2", level_type="prep", academic_year=year
        )
        short.wing = klass.wing
        short.save(update_fields=["wing"])
        _periods(school, klass, teacher, 3)
        _periods(school, short, UserFactory(national_id="29300000096"), 2, at_hour=12)

        register = wing_register(klass.wing, SUNDAY)

        assert register.width == 3
        assert [len(cells) for _section, cells in register.matrix] == [3, 3]
        assert register.matrix[1][1][-1] is None


# ══════════════════════════════════════════════════════════════════
# المخارج
# ══════════════════════════════════════════════════════════════════


class TestTheExports:
    def test_the_print_view_has_no_national_id(
        self, client_as, school, klass, kids, teacher, supervisor
    ):
        _confirm(klass, _periods(school, klass, teacher, 2)[0], {kids[0]: "absent"}, supervisor)

        body = client_as(supervisor).get(_section_url(klass)).content.decode()

        assert "كشفٌ جزئيٌّ حتى" in body
        assert "طالب 0" in body
        for kid in kids:
            assert kid.national_id not in body, "لا رقمَ شخصيّاً في الكشف"
        assert body.count('class="sig-item"') == 1
        assert "النائب الإداري" not in body and "مدير المدرسة" not in body
        assert 'class="emblem"' not in body, "كشفُ الحصص بلا شعار (طوارئ D-386م)"

    def test_the_excel_is_protected_and_has_no_national_id(
        self, client_as, school, klass, kids, teacher, supervisor
    ):
        _confirm(klass, _periods(school, klass, teacher, 2)[0], {}, supervisor)

        response = client_as(supervisor).get(_section_url(klass, "xlsx"))

        assert response.status_code == 200
        assert response["Content-Disposition"].isascii()
        book = openpyxl.load_workbook(io.BytesIO(response.content))
        sheet = book.active
        assert sheet.protection.sheet
        assert sheet.sheet_view.rightToLeft
        values = " ".join(str(c.value) for row in sheet.iter_rows() for c in row if c.value)
        assert "طالب 0" in values
        for kid in kids:
            assert kid.national_id not in values

    def test_the_wing_workbook_is_a_summary_then_a_sheet_per_section(
        self, client_as, school, klass, kids, teacher, supervisor
    ):
        _periods(school, klass, teacher, 2)

        response = client_as(supervisor).get(_wing_url(klass.wing.code, "xlsx"))

        book = openpyxl.load_workbook(io.BytesIO(response.content))
        assert book.sheetnames == ["ملخّص الجناح", klass.short_code]
        assert all(sheet.protection.sheet for sheet in book.worksheets)

    def test_the_pdf_downloads(self, client_as, school, klass, kids, teacher, supervisor):
        _periods(school, klass, teacher, 1)

        response = client_as(supervisor).get(_section_url(klass, "pdf"))

        assert response.status_code in (200, 503), "503 حين لا تتوفّر خدمةُ التوليد"
        if response.status_code == 200:
            assert response["Content-Type"] == "application/pdf"
            assert response["Content-Disposition"].startswith("attachment;")

    def test_the_wing_summary_shows_entered_periods_and_one_signature(
        self, client_as, school, klass, kids, teacher, supervisor
    ):
        periods = _periods(school, klass, teacher, 2)
        _enter(periods[0], kids[0], teacher, "absent")
        _enter(periods[0], kids[1], teacher, "present")

        body = client_as(supervisor).get(_wing_url(klass.wing.code)).content.decode()

        assert body.count('class="sig-item"') == 2, "توقيعٌ في ملخّص الجناح وآخرُ في كشف الشعبة"
        assert "غ 1 · م 0" in body, "رقمُ الغياب من الإدخال لا «لم تُرصد» وحدَها"

    def test_the_excel_has_one_signature_and_the_entered_numbers(
        self, client_as, school, klass, kids, teacher, supervisor
    ):
        periods = _periods(school, klass, teacher, 1)
        _enter(periods[0], kids[0], teacher, "absent")

        response = client_as(supervisor).get(_wing_url(klass.wing.code, "xlsx"))

        book = openpyxl.load_workbook(io.BytesIO(response.content))
        for sheet in book.worksheets:
            values = [str(c.value) for row in sheet.iter_rows() for c in row if c.value]
            # المسمّى والاسم: اسمُ المشرف في التجهيز يطابق مسمّاه، فتوقيعٌ واحدٌ = خليّتان.
            assert values.count("مشرف الجناح") == 2
            assert not any("النائب الإداري" in v or "مدير المدرسة" in v for v in values)
        section = book[klass.short_code]
        assert "رُصدت" in [str(c.value) for row in section.iter_rows() for c in row if c.value]


class TestTheOrientation:
    """أفقيٌّ افتراضاً، وعموديٌّ بالاختيار — صفحةٌ واحدةٌ لكلّ فصل (قرارُ 2026-09-13)."""

    def test_the_print_view_follows_the_chosen_orientation(
        self, client_as, school, klass, kids, teacher, supervisor
    ):
        _periods(school, klass, teacher, 2)
        client = client_as(supervisor)

        assert "size: A4 landscape" in client.get(_section_url(klass)).content.decode()
        portrait = client.get(_section_url(klass) + "&orient=portrait").content.decode()
        assert "size: A4 portrait" in portrait
        assert "page-break-inside: avoid" in portrait, "العموديُّ صفحةٌ واحدةٌ للفصل"

    def test_a_portrait_excel_fits_each_sheet_on_one_page(
        self, client_as, school, klass, kids, teacher, supervisor
    ):
        _periods(school, klass, teacher, 2)

        response = client_as(supervisor).get(
            _wing_url(klass.wing.code, "xlsx") + "&orient=portrait"
        )

        book = openpyxl.load_workbook(io.BytesIO(response.content))
        for sheet in book.worksheets:
            assert sheet.page_setup.orientation == "portrait"
            assert (sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight) == (1, 1)


class TestTheFooter:
    """الذيلُ في المخارج الثلاثة — كان في الـPDF وحدَه (ملاحظةُ 2026-09-13)."""

    def test_the_print_view_ends_each_class_page_with_the_footer(
        self, client_as, school, klass, kids, teacher, supervisor
    ):
        _periods(school, klass, teacher, 2)

        body = client_as(supervisor).get(_wing_url(klass.wing.code)).content.decode()

        assert body.count('class="sheet-footer"') == 2, "ملخّصُ الجناح وكشفُ الشعبة"
        assert "مُتَعَلِّمٌ رِيَادِيٌّ لِتَنْمِيَةٍ مُسْتَدَامَةٍ" in body

    def test_every_excel_sheet_ends_with_the_footer_inside_its_print_area(
        self, client_as, school, klass, kids, teacher, supervisor
    ):
        _periods(school, klass, teacher, 2)

        response = client_as(supervisor).get(_wing_url(klass.wing.code, "xlsx"))

        book = openpyxl.load_workbook(io.BytesIO(response.content))
        for sheet in book.worksheets:
            last = sheet.cell(row=sheet.max_row, column=1).value or ""
            assert "وزارة التربية والتعليم" in last
            assert sheet.print_area.endswith(f"${sheet.max_row}")


class TestOnlyHisOwnWing:
    def test_another_wings_supervisor_is_turned_away(
        self, client_as, school, year, klass, kids, teacher, supervisor
    ):
        other = UserFactory(full_name="مشرف جناح آخر", national_id="29300000093")
        MembershipFactory(
            user=other, school=school, role=RoleFactory(school=school, name="admin_supervisor")
        )
        Wing.objects.create(
            school=school, code="w2", name="جناح 2", academic_year=year, supervisor=other
        )

        assert client_as(other).get(_section_url(klass, "xlsx")).status_code == 404
        assert client_as(other).get(_wing_url(klass.wing.code, "xlsx")).status_code == 404
        assert client_as(other).get(_wing_url("w2", "xlsx")).status_code == 200

    def test_a_teacher_is_refused(self, client_as, school, klass, teacher, supervisor):
        response = client_as(teacher).get(_section_url(klass))

        assert response.status_code in (302, 403)


# ══════════════════════════════════════════════════════════════════
# علل الأدوات المشتركة
# ══════════════════════════════════════════════════════════════════


class TestTheSharedExportTools:
    def test_the_exporter_role_is_written_in_arabic(self, rf, school, supervisor):
        request = rf.get("/")
        request.user = supervisor

        context = get_export_context(request, "كشف")

        assert context["exporter_role"] == "مشرف إداري"
        assert context["exporter_role_code"] == "admin_supervisor"

    @pytest.mark.parametrize(
        ("filename", "fallback"),
        [("كشف_الجناح.xlsx", "document.xlsx"), ("register_11-5.xlsx", "register_11-5.xlsx")],
    )
    def test_an_excel_name_keeps_its_extension(self, filename, fallback):
        header = _content_disposition(filename, True)

        assert header.isascii()
        assert f'filename="{fallback}"' in header

    def test_the_printed_page_header_is_the_schools_own_name(self):
        from reports.services import ExcelService

        book, sheet, _styles = ExcelService._make_workbook("ورقة")
        ExcelService._add_professional_header(sheet, "مدرسة التجربة", "تقرير", "2026-2027", 3)
        ExcelService._setup_print(sheet, 3, 1)

        assert sheet.oddHeader.center.text.endswith("مدرسة التجربة")
