"""شعارُ المنصّة «الكامل بالكتاب» على الشاشة وشعارُ الوزارة للمخرجات وحدَها (W-20261014-009، قرار المالك 10-14؛ يحلّ محلّ «ش» v3).

شاشةً: `static/brand/schoolos-full.svg` على شريحةٍ داكنةٍ محايدة (CSS) لا على العنّابيّ. مخرجاتٍ (PDF وطباعة وExcel): شعارُ الوزارة
الأصليّ أعلى وسط الترويسة واسمُ المدرسة تحته، بلا سطرِ نصّ الوزارة. وحارسُ «لا شعارَ وزارةٍ في المنصّة» في `test_ministry_logo_exports_only.py`.
"""

import io
import pathlib
import re

import pytest
from django.template import Context, Template
from django.template.loader import render_to_string

from core import brand, ministry_logo, school_logo
from tests.test_observation_pdf_form import criteria, observation  # noqa: F401 — تجهيزاتٌ مشتركة

ROOT = pathlib.Path(__file__).resolve().parent.parent
EMBLEM = re.compile(r"emblem(?:-white)?\.svg|logoMaroon|logowhite")


def _luminance(hex_colour: str) -> float:
    channels = [int(hex_colour.lstrip("#")[i : i + 2], 16) / 255 for i in (0, 2, 4)]
    lin = [c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4 for c in channels]
    return 0.2126 * lin[0] + 0.7152 * lin[1] + 0.0722 * lin[2]


def _contrast(a: str, b: str) -> float:
    high, low = sorted((_luminance(a), _luminance(b)), reverse=True)
    return (high + 0.05) / (low + 0.05)


def _blend(over: str, under: str, alpha: float) -> str:
    over_rgb = [int(over.lstrip("#")[i : i + 2], 16) for i in (0, 2, 4)]
    under_rgb = [int(under.lstrip("#")[i : i + 2], 16) for i in (0, 2, 4)]
    mixed = [round(o * alpha + u * (1 - alpha)) for o, u in zip(over_rgb, under_rgb)]
    return "#" + "".join(f"{c:02x}" for c in mixed)


# ── الأصلُ والوسم ─────────────────────────────────────────────────────────


def test_the_source_is_the_full_logo_svg_with_a_viewbox():
    svg = (ROOT / "static/brand/schoolos-full.svg").read_text(encoding="utf-8")
    assert school_logo.LOGO_STATIC_PATH == "brand/schoolos-full.svg"
    assert re.search(r'viewBox="0 0 [\d.]+ [\d.]+"', svg)


def test_the_tag_gives_the_static_link_on_screen():
    on_screen = Template('{% load school_logo %}{% school_logo "nav" %}').render(Context())
    assert 'class="school-logo school-logo--nav"' in on_screen
    assert "brand/schoolos-full.svg" in on_screen
    assert 'alt=""' in on_screen  # زخرفيٌّ: اسمُ المدرسة نصٌّ يجاوره
    assert "data:" not in on_screen


def test_the_tag_refuses_an_unknown_size():
    with pytest.raises(Exception, match="حجمٌ غير معروف"):
        Template('{% load school_logo %}{% school_logo "huge" %}').render(Context())


# ── الشاشة ────────────────────────────────────────────────────────────────


@pytest.mark.django_db
def test_the_login_page_shows_the_logo_above_the_school_name(client):
    from django.urls import reverse

    html = client.get(reverse("login"), follow=True).content.decode()
    assert "school-logo--login" in html
    assert html.index("school-logo--login") < html.index("login-school-name")


def test_the_offline_and_error_pages_carry_the_logo():
    assert "school-logo--login" in render_to_string("pwa/offline_global.html", {})
    error = render_to_string(
        "errors/_error_page.html",
        {
            "code": "404",
            "heading": "x",
            "desc": "x",
            "icon": "search",
            "accent": "MAROON",
            "accent_dark": "MAROON_FG",
            "login_header": True,
        },
    )
    assert "school-logo--error" in error


def test_the_platform_header_and_admin_header_carry_the_logo_before_the_school_name():
    base = (ROOT / "templates/base/base.html").read_text(encoding="utf-8")
    nav = base.split('class="nb nav-logo"', 1)[1].split("</a>", 1)[0]
    assert '{% school_logo "nav" %}' in nav and nav.index("school_logo") < nav.index(
        "nav-logo-text"
    )
    admin = (ROOT / "templates/admin/base_site.html").read_text(encoding="utf-8")
    assert '{% school_logo "nav" %}' in admin.split('class="adm-brand"', 1)[1].split("</a>", 1)[0]


# ── Excel ─────────────────────────────────────────────────────────────────


def test_every_excel_title_carries_the_ministry_logo_pinned_in_a1_without_adding_rows():
    from openpyxl import Workbook, load_workbook

    from core.export_utils import add_excel_title_rows

    wb = Workbook()
    ws = wb.active
    add_excel_title_rows(ws, 6, "مدرسةٌ", "تقرير", "صدر")
    assert ws.max_row == 3, "الشعارُ صورةٌ لا صفوف"
    assert len(ws._images) == 1
    image = ws._images[0]
    assert image.anchor == "A1"
    assert image.width == image.height == ministry_logo.EXCEL_SIDE_PX
    assert ministry_logo.EXCEL_SIDE_PX <= 40 / 0.75, "يسكن ارتفاعَ الصفّ الأوّل (40pt)"

    buffer = io.BytesIO()
    wb.save(buffer)
    reloaded = load_workbook(io.BytesIO(buffer.getvalue()))
    assert len(reloaded.active._images) == 1  # بقي في الملفّ المحفوظ لا في الذاكرة وحدَها


# ── PDF (مقيساً لا مقروءاً من القالب) ─────────────────────────────────────


def _fonts_of(pdf: bytes) -> set[str]:
    from pypdf import PdfReader

    page = PdfReader(io.BytesIO(pdf)).pages[0]
    return {str(font.get_object()["/BaseFont"]) for font in page["/Resources"]["/Font"].values()}


@pytest.mark.django_db
def test_the_observation_pdf_embeds_the_logo_without_changing_its_page_count(observation):  # noqa: F811
    """استمارةُ الزيارة بخطّ نموذج المدرسة الأصليّ (Traditional Arabic ← Noto Naskh) بقرارٍ سابق؛ فيُقاس هنا الشعارُ وأنّه لا يغيّر عددَ الصفحات."""
    from pypdf import PdfReader

    from core.pdf_utils import render_pdf_bytes
    from quality.observation_views import _pdf_context

    html = render_to_string("quality/observation_pdf.html", _pdf_context(observation))
    assert "data:image/png;base64," in html and 'class="ministry-logo"' in html
    reader = PdfReader(io.BytesIO(render_pdf_bytes(html)))
    assert len(reader.pages[0].images) >= 1, "الشعارُ صورةٌ في الصفحة الأولى"
    bare = re.sub(r'<img class="ministry-logo[^>]*>', "", html)
    assert len(reader.pages) == len(
        PdfReader(io.BytesIO(render_pdf_bytes(bare))).pages
    ), "الشعارُ لم يدفع شيئاً إلى صفحةٍ جديدة"


@pytest.mark.django_db
def test_the_platform_report_pdf_carries_the_logo_and_is_set_in_tajawal():
    """مقيساً من الملفّ المولَّد بمسار المنصّة الحقيقيّ (`pdf_utils` يحقن Tajawal بمسارات مطلقة) لا من نصّ القالب."""
    from django.template import engines
    from pypdf import PdfReader

    from core.pdf_utils import render_pdf_bytes

    child = engines["django"].from_string(
        '{% extends "reports/base_qatar_report.html" %}{% block school_name %}مدرسةٌ تجريبيّة{% endblock %}'
    )
    html = child.render({"school": type("S", (), {"name": "مدرسةٌ تجريبيّة", "vision": "رؤية"})()})
    pdf = render_pdf_bytes(html)
    assert len(PdfReader(io.BytesIO(pdf)).pages[0].images) >= 1
    fonts = _fonts_of(pdf)
    assert any(
        "Tajawal" in name for name in fonts
    ), f"خطُّ الـPDF يجب أن يكون Tajawal لا بديلاً: {fonts}"
    assert not any("DejaVu" in name for name in fonts), fonts


# ── التباين ───────────────────────────────────────────────────────────────


def test_the_logo_chip_is_a_neutral_dark_surface_not_maroon():
    """الشريحةُ `DARK["SURFACE"]` محايدةٌ (الزرقةُ أعلى من الحمرة) وأغمقُ من العنّابيّ فتبرز عليه."""
    chip = brand.DARK["SURFACE"]
    red, _, blue = (int(chip.lstrip("#")[i : i + 2], 16) for i in (0, 2, 4))
    assert blue > red, chip
    assert _luminance(chip) < _luminance(brand.MAROON) + 0.05


# ── الغياب: لا شعارَ وزارة في أيّ مصدر ────────────────────────────────────


def test_no_source_refers_to_the_ministry_emblem_images():
    skip = {
        ".git",
        ".claude",
        "node_modules",
        "staticfiles",
        "tests",
        "docs",
        "AAdocs",
        "migrations",
        "venv",
        ".venv",
    }
    offenders = []
    for path in ROOT.rglob("*"):
        if path.suffix not in {".py", ".html", ".js", ".css", ".txt"} or not path.is_file():
            continue
        if skip & set(path.relative_to(ROOT).parts):
            continue
        if EMBLEM.search(path.read_text(encoding="utf-8", errors="ignore")):
            offenders.append(path.relative_to(ROOT).as_posix())
    assert offenders == [], f"مراجعُ لصور شعار الوزارة: {offenders}"
