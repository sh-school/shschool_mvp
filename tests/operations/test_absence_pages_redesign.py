"""إعادةُ تصميم صفحات متابعة الغياب (W-20261020-005، الخطوتان 1 و2): القائمةُ بثلاثة بنودٍ مباشرة، والترويسةُ والبحثُ موحَّدان."""

import pytest
from django.urls import reverse

from tests.operations.test_absence_followup import (  # noqa: F401
    DAY,
    _calendar_first,
    _get,
    absent_both,
    second_wing,
    staff,
)
from tests.test_absence_excuses import MONDAY, _absent_day  # noqa: F401
from tests.test_period_register import (  # noqa: F401
    SUNDAY,
    kids,
    klass,
    other_teacher,
    subjects,
    supervisor,
    teacher,
    year,
)

pytestmark = pytest.mark.django_db


def test_the_menu_has_three_direct_items_for_absence_follow_up(client_as, staff, absent_both):
    body = _get(client_as, staff["principal"]).content.decode()
    for name in ("wings:ministry_report", "daily_report", "absence_actions"):
        assert f'href="{reverse(name)}"' in body, name


@pytest.mark.parametrize("name", ["daily_report", "absence_actions"])
def test_the_follow_up_pages_share_the_unified_filter_bar(client_as, staff, absent_both, name):
    body = _get(client_as, staff["principal"], name).content.decode()
    assert 'role="search"' in body
    for field in ('name="date"', 'name="wing"', 'name="grade"', 'name="q"'):
        assert field in body, field


def test_the_search_filters_by_student_name_and_is_carried_to_the_print_link(
    client_as, staff, absent_both
):
    body = _get(client_as, staff["principal"], q="الجناح الثاني").content.decode()
    assert "طالب الجناح الثاني" in body and "طالب 0" not in body
    assert "q=" in body  # يُحمل الترشيحُ إلى PDF والعودة


def test_an_unmatched_search_shows_the_empty_state(client_as, staff, absent_both):
    body = _get(client_as, staff["principal"], q="لا أحد بهذا الاسم").content.decode()
    assert "لا غائبَ يطابق التصفية" in body


def test_the_ministry_page_has_the_wing_filter_for_the_wide_role(client_as, staff, absent_both):
    resp = client_as(staff["principal"]).get(reverse("wings:ministry_report"), {"date": DAY})
    assert resp.status_code == 200
    assert 'name="wing"' in resp.content.decode()


def test_the_export_buttons_are_pdf_and_excel_only(client_as, staff, absent_both):
    body = _get(client_as, staff["principal"]).content.decode()
    assert "تصدير PDF" in body and "تصدير Excel" in body
    assert reverse("absence_followup_xlsx") in body


def test_excel_export_matches_the_filtered_screen(client_as, staff, absent_both):
    import io

    from openpyxl import load_workbook

    resp = _get(
        client_as, staff["principal"], "absence_followup_xlsx", page="contact", q="الجناح الثاني"
    )
    assert resp.status_code == 200 and "spreadsheetml" in resp["Content-Type"]
    ws = load_workbook(io.BytesIO(resp.content)).active
    cells = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "طالب الجناح الثاني" in cells and "طالب 0" not in cells
    assert "أصدر الملف" in cells


def test_excel_export_is_denied_to_a_teacher(client_as, teacher, absent_both):
    assert _get(client_as, teacher, "absence_followup_xlsx").status_code == 403


def test_pdf_footer_names_the_issuer_and_has_a_class_column(staff, school, absent_both):
    from operations.absence_followup_print import build_followup_pdf

    user = staff["principal"]
    pdf = build_followup_pdf(school, user, {"date": DAY, "page": "contact"}).content
    import io
    import unicodedata

    import pdfplumber

    text = " ".join(p.extract_text() or "" for p in pdfplumber.open(io.BytesIO(pdf)).pages)
    flat = unicodedata.normalize("NFKC", text)
    both = flat + " " + flat[::-1]
    assert unicodedata.normalize("NFKC", "أصدر الملف") in both
    assert unicodedata.normalize("NFKC", "الصف") in both
