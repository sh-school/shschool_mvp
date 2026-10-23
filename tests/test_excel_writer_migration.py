"""ترحيلُ مولّدات Excel إلى الكاتب المركزيّ (W-20261021-004، CWE-1236).

وحدويّ: ``safe_cell``/``safe_set``/``safe_append`` تحيِّد النصَّ وتترك الأرقامَ والتواريخ.
تكامليّ: نصٌّ خبيثٌ يبدأ بـ``=`` أو ``@`` في اسم طالب/معلّم أو وصف مخالفة يخرج من كلّ
مولّدٍ رُحِّل نصّاً مسبوقاً بـ``'`` لا صيغةً حيّة.
"""

from __future__ import annotations

import datetime as dt
import io
from decimal import Decimal

import openpyxl
import pytest
from django.urls import reverse

from core.excel_safety import safe_append, safe_cell, safe_set, safe_value
from tests.conftest import BehaviorInfractionFactory
from tests.test_period_register import (  # noqa: F401 — تجهيزاتُ الكشف نفسُها
    SUNDAY,
    kids,
    klass,
    supervisor,
    teacher,
    year,
)
from tests.test_schedule_export import principal, teaching_school  # noqa: F401

pytestmark = pytest.mark.django_db

EVIL_EQ = '=HYPERLINK("http://evil","x")'
EVIL_AT = "@SUM(1+1)"


def _texts(book) -> list[str]:
    return [
        c.value
        for ws in book.worksheets
        for row in ws.iter_rows()
        for c in row
        if isinstance(c.value, str)
    ]


def _assert_no_live_formula(texts: list[str], needle: str) -> None:
    hits = [t for t in texts if needle in t]
    assert hits, f"النصُّ الخبيث يجب أن يظهر مُحيَّداً لا محذوفاً: {needle}"
    assert all(t.startswith("'") for t in hits), hits


def _rename(user, name: str) -> None:
    user.full_name = name
    user.save(update_fields=["full_name"])


# ══════════════════════════════════════════════════════════════════
# الكاتب المركزيّ — وحدويّ
# ══════════════════════════════════════════════════════════════════


class TestCentralWriter:
    def test_safe_cell_neutralizes_text_and_returns_the_cell(self):
        ws = openpyxl.Workbook().active
        cell = safe_cell(ws, 1, 1, EVIL_EQ)
        assert cell.value == "'" + EVIL_EQ
        assert cell.data_type == "s"

    @pytest.mark.parametrize("prefix", ["=", "+", "-", "@"])
    def test_every_formula_prefix_is_neutralized(self, prefix):
        ws = openpyxl.Workbook().active
        assert safe_cell(ws, row=2, column=3, value=f"{prefix}1+1").value == f"'{prefix}1+1"

    def test_numbers_dates_bool_and_none_pass_through_with_their_type(self):
        ws = openpyxl.Workbook().active
        today = dt.date(2026, 10, 21)
        assert safe_cell(ws, 1, 1, -7).value == -7
        assert ws.cell(row=1, column=1).data_type == "n"
        assert safe_cell(ws, 2, 1, 3.5).value == 3.5
        assert safe_cell(ws, 3, 1, Decimal("2.50")).value == Decimal("2.50")
        assert safe_cell(ws, 4, 1, today).value == today
        assert safe_cell(ws, 5, 1, True).value is True
        assert safe_cell(ws, 6, 1, None).value is None

    def test_safe_cell_without_value_reads_or_formats_without_overwriting(self):
        ws = openpyxl.Workbook().active
        safe_cell(ws, 1, 1, "حاضر")
        assert safe_cell(ws, 1, 1).value == "حاضر"

    def test_an_ordinary_arabic_name_is_unchanged(self):
        ws = openpyxl.Workbook().active
        assert safe_cell(ws, 1, 1, "سفيان مسيف").value == "سفيان مسيف"

    def test_safe_set_neutralizes_cell_value(self):
        ws = openpyxl.Workbook().active
        cell = ws.cell(row=1, column=1)
        safe_set(cell, EVIL_AT)
        assert cell.value == "'" + EVIL_AT

    def test_safe_append_neutralizes_every_value_of_the_row(self):
        ws = openpyxl.Workbook().active
        safe_append(ws, [EVIL_EQ, 5, None, "عادي"])
        values = [c.value for c in ws[1]]
        assert values[0] == "'" + EVIL_EQ
        assert values[1] == 5
        assert values[3] == "عادي"

    def test_safe_append_accepts_a_generator_and_a_dict(self):
        ws = openpyxl.Workbook().active
        safe_append(ws, (v for v in (EVIL_AT, 1)))
        safe_append(ws, {"A": EVIL_EQ, "B": 2})
        assert ws["A1"].value == "'" + EVIL_AT
        assert ws["A2"].value == "'" + EVIL_EQ and ws["B2"].value == 2

    def test_safe_value_is_idempotent(self):
        once = safe_value(EVIL_EQ)
        assert safe_value(once) == once


# ══════════════════════════════════════════════════════════════════
# المولّدات المرحَّلة — تكامليّ
# ══════════════════════════════════════════════════════════════════


class TestRegisterWorkbook:
    def test_a_student_name_starting_with_equals_is_not_a_formula(
        self, client_as, school, klass, kids, teacher, supervisor
    ):
        _rename(kids[0], EVIL_EQ)
        url = reverse("wings:section_register", args=[klass.id])
        resp = client_as(supervisor).get(f"{url}?date={SUNDAY.isoformat()}&format=xlsx")
        assert resp.status_code == 200
        book = openpyxl.load_workbook(io.BytesIO(resp.content))
        _assert_no_live_formula(_texts(book), "HYPERLINK")

    def test_the_wing_workbook_is_neutralized_too(
        self, client_as, school, klass, kids, teacher, supervisor
    ):
        _rename(kids[1], EVIL_AT)
        url = reverse("wings:wing_register", args=[klass.wing.code])
        resp = client_as(supervisor).get(f"{url}?date={SUNDAY.isoformat()}&format=xlsx")
        book = openpyxl.load_workbook(io.BytesIO(resp.content))
        _assert_no_live_formula(_texts(book), "SUM(1+1)")


class TestScheduleExport:
    def test_a_teacher_name_starting_with_equals_is_not_a_formula(self, principal, teaching_school):
        from django.test import Client

        _rename(teaching_school[0], EVIL_EQ)
        client = Client()
        client.force_login(principal)
        resp = client.get(
            reverse("schedule_export_excel") + "?view=all_teachers&paper=a3",
            HTTP_HOST="localhost",
            follow=True,
        )
        assert resp.status_code == 200
        book = openpyxl.load_workbook(io.BytesIO(resp.content))
        _assert_no_live_formula(_texts(book), "HYPERLINK")


class TestAcademicReportsExcel:
    def test_build_neutralizes_free_text_columns(self, school):
        from reports.services import AcademicReportsExcel

        book = AcademicReportsExcel._build(
            "ورقة", "عنوان", [("الاسم", 20), ("ملاحظة", 20)], [[EVIL_EQ, EVIL_AT]], school
        )
        texts = _texts(book)
        _assert_no_live_formula(texts, "HYPERLINK")
        _assert_no_live_formula(texts, "SUM(1+1)")


class TestExcelServiceReports:
    def test_class_attendance_report_neutralizes_student_names(self, school, klass, kids):
        from reports.services import ExcelService

        _rename(kids[0], EVIL_EQ)
        resp = ExcelService.attendance_excel(klass, school)
        book = openpyxl.load_workbook(io.BytesIO(resp.content))
        _assert_no_live_formula(_texts(book), "HYPERLINK")

    def test_class_results_report_neutralizes_student_names(self, school, klass, kids):
        from reports.services import ExcelService

        _rename(kids[0], EVIL_EQ)
        resp = ExcelService.class_results_excel(klass, school)
        book = openpyxl.load_workbook(io.BytesIO(resp.content))
        _assert_no_live_formula(_texts(book), "HYPERLINK")

    def test_behavior_report_neutralizes_the_description(self, school, kids, teacher):
        from reports.services import ExcelService

        BehaviorInfractionFactory(
            school=school, student=kids[0], reported_by=teacher, description=EVIL_AT
        )
        resp = ExcelService.behavior_excel(school)
        book = openpyxl.load_workbook(io.BytesIO(resp.content))
        _assert_no_live_formula(_texts(book), "SUM(1+1)")


class TestStudentExports:
    @staticmethod
    def _student(school, name, national_id):
        from tests.conftest import (
            ClassGroupFactory,
            MembershipFactory,
            RoleFactory,
            StudentEnrollmentFactory,
            UserFactory,
        )

        klass = ClassGroupFactory(school=school)
        student = UserFactory(full_name=name, national_id=national_id)
        MembershipFactory(
            user=student, school=school, role=RoleFactory(school=school, name="student")
        )
        StudentEnrollmentFactory(student=student, class_group=klass)
        return student

    def test_core_student_export_neutralizes_the_name(self, client_as, principal_user, school):
        self._student(school, EVIL_EQ, "99900000701")
        resp = client_as(principal_user).get(reverse("student_export_excel"))
        assert resp.status_code == 200
        book = openpyxl.load_workbook(io.BytesIO(resp.content))
        _assert_no_live_formula(_texts(book), "HYPERLINK")

    def test_student_affairs_export_neutralizes_the_name(
        self, client_as, principal_user, school, seeded_calendar
    ):
        from tests.conftest import ClassGroupFactory, StudentEnrollmentFactory

        student = self._student(school, EVIL_AT, "99900000702")
        klass = ClassGroupFactory(school=school, academic_year=seeded_calendar)
        StudentEnrollmentFactory(student=student, class_group=klass)
        resp = client_as(principal_user).get(reverse("student_affairs:student_export"))
        assert resp.status_code == 200
        book = openpyxl.load_workbook(io.BytesIO(resp.content))
        _assert_no_live_formula(_texts(book), "SUM(1+1)")
