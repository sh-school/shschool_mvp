"""سجلُّ القيد بشكليه — أوراقُ الصفوف، والورقةُ المسطّحة — والحزمةُ من stdin."""

import io

import pytest
from django.core.management import call_command
from django.core.management.base import CommandError

from core.management.commands.import_enrolment_register import Command, _pack, _unpack

pytestmark = pytest.mark.django_db

ROWS = [
    ("31400000001", "أحمد الأوّل", "7", "07/1"),
    ("31400000002", "سالم الثاني", "7", "07/2"),
    ("31400000003", "خالد الثالث", "8", "08/1"),
]


def _flat(tmp_path):
    """تصديرُ 2026-2027: ورقةٌ واحدة «الجميع» ورؤوسُها في السطر الأوّل."""
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "الجميع"
    sheet.append(["الرقم الشخصي", "الاسم", "الصف ", "الصف/الشعبة"])
    for row in ROWS:
        sheet.append(list(row))
    path = tmp_path / "مسطّح.xlsx"
    book.save(path)
    return str(path)


def _by_grade(tmp_path):
    """تصديرُ 2025-2026: ورقةٌ لكلّ صفّ ورؤوسُها في السطر الثالث."""
    import openpyxl

    book = openpyxl.Workbook()
    book.remove(book.active)
    for grade in ("07", "08"):
        sheet = book.create_sheet(f"الصف- {grade}")
        sheet.append(["سجلّ القيد", "", "", ""])
        sheet.append(["", "", "", ""])
        sheet.append(["الرقم الشخصي", "الاسم", "الصف", "الصف/الشعبة"])
        for row in ROWS:
            if row[2] == grade.lstrip("0"):
                sheet.append(list(row))
    path = tmp_path / "بالصفوف.xlsx"
    book.save(path)
    return str(path)


class TestBothShapes:
    def test_the_flat_export_is_read(self, tmp_path):
        roster, _tracks = Command()._read(_flat(tmp_path))
        assert len(roster) == 3
        assert roster["31400000001"] == ("أحمد الأوّل", "7/1")

    def test_the_old_export_is_still_read(self, tmp_path):
        """الملفُّ القديم لا يُكسر بتعميم القارئ."""
        roster, _tracks = Command()._read(_by_grade(tmp_path))
        assert len(roster) == 3
        assert roster["31400000003"] == ("خالد الثالث", "8/1")

    def test_a_sheet_without_headers_is_skipped_not_guessed(self, tmp_path):
        import openpyxl

        book = openpyxl.Workbook()
        book.active.title = "ملاحظات"
        book.active.append(["كلامٌ حرّ", "لا رؤوسَ له"])
        path = tmp_path / "بلا.xlsx"
        book.save(path)

        with pytest.raises(CommandError, match="لم يُقرأ طالبٌ واحد"):
            Command()._read(str(path))


class TestPackage:
    def test_a_round_trip_keeps_the_register_intact(self, tmp_path):
        roster, tracks = Command()._read(_flat(tmp_path))
        again, again_tracks = _unpack(_pack(roster, tracks))
        assert again == roster
        assert again_tracks == tracks

    def test_a_broken_package_says_so(self):
        with pytest.raises(CommandError, match="gzip"):
            _unpack("ليست حزمة")

    def test_an_empty_package_is_refused(self):
        with pytest.raises(CommandError, match="خالية"):
            _unpack(_pack({}, {}))

    def test_a_dash_reads_the_package_from_stdin(self, tmp_path, school, monkeypatch, capsys):
        """الحزمةُ في سطر الأوامر تُقرأ بـ`ps` وتُسجَّل — فالأنبوبُ هو الأسلم."""
        roster, tracks = Command()._read(_flat(tmp_path))
        monkeypatch.setattr("sys.stdin", io.StringIO(_pack(roster, tracks)))

        call_command("import_enrolment_register", rows_b64="-", year="2026-2027")
        assert "في السجلّ: 3 طالباً" in capsys.readouterr().out

    def test_a_path_and_a_package_together_are_refused(self, tmp_path, school):
        with pytest.raises(CommandError, match="واحداً منهما"):
            call_command(
                "import_enrolment_register", _flat(tmp_path), rows_b64="x", year="2026-2027"
            )

    def test_neither_a_path_nor_a_package_is_refused(self, school):
        with pytest.raises(CommandError, match="واحداً منهما"):
            call_command("import_enrolment_register", year="2026-2027")
